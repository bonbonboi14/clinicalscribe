from __future__ import annotations

import html
import io
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import UUID

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from models.clerking_sheet import ClerkingSheet
from models.differential import DifferentialResult
from models.treatment_plan import TreatmentPlan
from storage.database import Database


class ApprovalRequiredError(ValueError):
    """Raised when an export is requested before immutable clinician approval."""


@dataclass(frozen=True)
class RenderedExport:
    content: bytes
    media_type: str
    filename: str


@dataclass(frozen=True)
class ApprovedSessionSnapshot:
    session: dict[str, Any]
    note: dict[str, Any]
    clerking_sheet: ClerkingSheet
    treatment_plan: TreatmentPlan | None
    differential: DifferentialResult | None
    transcript: dict[str, Any] | None
    validation_findings: list[dict[str, Any]]
    audit_events: list[dict[str, Any]]

    @property
    def session_id(self) -> str:
        return self.session["id"]

    @property
    def specialty(self) -> str:
        return str(self.note.get("template_name") or "Clinical").replace("_", " ").title()

    @property
    def session_date(self) -> str:
        return str(self.session["created_at"])[:10]

    def metadata(self) -> dict[str, Any]:
        return {
            "session_id": self.session_id,
            "session_date": self.session_date,
            "specialty": self.specialty,
            "source_language": self.session.get("source_language"),
            "approved_note": {
                "id": self.note["id"],
                "version": self.note["version"],
                "approved_by": self.note.get("approved_by"),
                "approved_at": self.note.get("approved_at"),
                "template_name": self.note.get("template_name"),
                "engine": self.note.get("engine"),
            },
            "differential_included": bool(self.differential and self.differential.enabled),
            "audio_included": False,
        }

    def as_json(self) -> dict[str, Any]:
        return {
            "metadata": self.metadata(),
            "clerking_sheet": self.clerking_sheet.model_dump(mode="json"),
            "clinical_note": {
                key: value for key, value in self.note.items()
                if key not in {"plain_text"}
            },
            "treatment_plan": (
                self.treatment_plan.model_dump(mode="json") if self.treatment_plan else None
            ),
            "differential_diagnosis": (
                self.differential.model_dump(mode="json") if self.differential else None
            ),
            "claim_validation": self.validation_findings,
        }


def _rows(rows: list[Any]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]


def _markdown_value(value: Any, indent: int = 0) -> list[str]:
    prefix = "  " * indent
    if isinstance(value, dict):
        lines: list[str] = []
        for key, item in value.items():
            label = key.replace("_", " ").title()
            if isinstance(item, (dict, list)):
                lines.append(f"{prefix}- **{label}:**")
                lines.extend(_markdown_value(item, indent + 1))
            else:
                lines.append(f"{prefix}- **{label}:** {item}")
        return lines
    if isinstance(value, list):
        lines = []
        for item in value:
            if isinstance(item, (dict, list)):
                lines.append(f"{prefix}-")
                lines.extend(_markdown_value(item, indent + 1))
            else:
                lines.append(f"{prefix}- {item}")
        return lines
    return [f"{prefix}- {value}"]


def clerking_markdown(sheet: ClerkingSheet) -> str:
    excluded = {"id", "session_id", "version", "fact_ids", "evidence_by_field", "status", "created_at"}
    lines = ["# Medical Clerking Sheet", ""]
    for key, value in sheet.model_dump(mode="json").items():
        if key in excluded:
            continue
        lines.extend((f"## {key.replace('_', ' ').title()}", ""))
        if isinstance(value, (dict, list)):
            lines.extend(_markdown_value(value))
        else:
            lines.append(str(value))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _speaker_labelled_text(path: str | None) -> str:
    if not path:
        return "Speaker-labelled transcript not available.\n"
    source = Path(path)
    if not source.is_file():
        return "Speaker-labelled transcript not available.\n"
    payload = json.loads(source.read_text(encoding="utf-8"))
    lines = []
    for segment in payload.get("segments", []):
        speaker = (
            segment.get("speaker_display_name")
            or segment.get("speaker_role")
            or segment.get("diarization_label")
            or "UNKNOWN"
        )
        start = int(segment.get("start_ms", 0)) / 1000
        end = int(segment.get("end_ms", 0)) / 1000
        text = segment.get("original_text", "")
        lines.append(f"[{start:.1f}-{end:.1f}] {speaker}: {text}")
    return "\n".join(lines).rstrip() + "\n"


class ExportService:
    SUPPORTED_FORMATS = {"txt", "md", "docx", "pdf", "print", "json"}

    def __init__(self, database: Database) -> None:
        self.database = database

    def load_approved_snapshot(self, session_id: UUID | str) -> ApprovedSessionSnapshot:
        session_key = str(UUID(str(session_id)))
        with self.database.connect() as connection:
            session = connection.execute(
                "SELECT * FROM sessions WHERE id = ?", (session_key,)
            ).fetchone()
            note = connection.execute(
                "SELECT * FROM notes WHERE session_id = ? AND status = 'APPROVED' "
                "ORDER BY version DESC LIMIT 1", (session_key,),
            ).fetchone()
            if note is None:
                raise ApprovalRequiredError("clinician approval is required before export")
            sheet_row = connection.execute(
                "SELECT * FROM clerking_sheets WHERE id = ?", (note["clerking_sheet_id"],)
            ).fetchone()
            treatment_row = connection.execute(
                "SELECT * FROM treatment_plans WHERE session_id = ? ORDER BY version DESC LIMIT 1",
                (session_key,),
            ).fetchone()
            differential_row = connection.execute(
                "SELECT * FROM differentials WHERE session_id = ? ORDER BY version DESC LIMIT 1",
                (session_key,),
            ).fetchone()
            transcript = connection.execute(
                "SELECT * FROM transcripts WHERE session_id = ? ORDER BY version DESC LIMIT 1",
                (session_key,),
            ).fetchone()
            findings = connection.execute(
                "SELECT claim_text, state, evidence_json, explanation FROM validation_findings "
                "WHERE artefact_type = 'CLINICAL_NOTE' AND artefact_id = ? ORDER BY created_at, id",
                (note["id"],),
            ).fetchall()
            audit = connection.execute(
                "SELECT id, artefact_type, artefact_id, action, actor, before_json, after_json, "
                "metadata_json, created_at FROM audit_events WHERE session_id = ? ORDER BY id",
                (session_key,),
            ).fetchall()
            labelled = connection.execute(
                "SELECT storage_path FROM transcript_artifacts WHERE session_id = ? "
                "AND kind = 'SPEAKER_LABELLED_TRANSCRIPT' ORDER BY version DESC LIMIT 1",
                (session_key,),
            ).fetchone()
        if session is None or sheet_row is None:
            raise ValueError("approved note has incomplete session provenance")
        note_data = dict(note)
        transcript_data = dict(transcript) if transcript else None
        if transcript_data is not None:
            transcript_data["speaker_labelled_path"] = labelled["storage_path"] if labelled else None
        return ApprovedSessionSnapshot(
            session=dict(session),
            note=note_data,
            clerking_sheet=ClerkingSheet.model_validate_json(sheet_row["structured_json"]),
            treatment_plan=(
                TreatmentPlan.model_validate_json(treatment_row["structured_json"])
                if treatment_row and treatment_row["structured_json"] != "{}" else None
            ),
            differential=(
                DifferentialResult.model_validate_json(differential_row["result_json"])
                if differential_row and differential_row["result_json"] != "{}" else None
            ),
            transcript=transcript_data,
            validation_findings=[
                {
                    "claim_text": row["claim_text"],
                    "state": row["state"],
                    "evidence_segment_ids": json.loads(row["evidence_json"]),
                    "explanation": row["explanation"],
                }
                for row in findings
            ],
            audit_events=[
                {
                    **dict(row),
                    "before": json.loads(row["before_json"]) if row["before_json"] else None,
                    "after": json.loads(row["after_json"]) if row["after_json"] else None,
                    "metadata": json.loads(row["metadata_json"] or "{}"),
                }
                for row in audit
            ],
        )

    def render(self, snapshot: ApprovedSessionSnapshot, export_format: str) -> RenderedExport:
        export_format = export_format.lower()
        if export_format not in self.SUPPORTED_FORMATS:
            raise ValueError(f"unsupported export format: {export_format}")
        stem = f"clinical-note-{snapshot.session_id}"
        if export_format == "txt":
            return RenderedExport(snapshot.note["plain_text"].encode("utf-8"), "text/plain; charset=utf-8", f"{stem}.txt")
        if export_format == "md":
            return RenderedExport(snapshot.note["content"].encode("utf-8"), "text/markdown; charset=utf-8", f"{stem}.md")
        if export_format == "json":
            data = json.dumps(snapshot.as_json(), ensure_ascii=False, indent=2).encode("utf-8")
            return RenderedExport(data, "application/json", f"{stem}.json")
        if export_format == "docx":
            return RenderedExport(self._docx(snapshot), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{stem}.docx")
        if export_format == "pdf":
            return RenderedExport(self._pdf(snapshot), "application/pdf", f"{stem}.pdf")
        return RenderedExport(self._print_html(snapshot).encode("utf-8"), "text/html; charset=utf-8", f"{stem}-print.html")

    def github_files(self, snapshot: ApprovedSessionSnapshot, *, push_transcripts: bool) -> dict[str, bytes]:
        files = {
            "metadata.json": json.dumps(snapshot.metadata(), ensure_ascii=False, indent=2).encode("utf-8"),
            "clerking_sheet.md": clerking_markdown(snapshot.clerking_sheet).encode("utf-8"),
            "clinical_note.md": snapshot.note["content"].encode("utf-8"),
            "treatment_plan.md": (
                (snapshot.treatment_plan.render_text() + "\n").encode("utf-8")
                if snapshot.treatment_plan else b"Treatment plan not generated.\n"
            ),
            "audit_log.json": json.dumps(snapshot.audit_events, ensure_ascii=False, indent=2).encode("utf-8"),
        }
        if snapshot.differential and snapshot.differential.enabled:
            files["differential_diagnosis.md"] = (snapshot.differential.render_text() + "\n").encode("utf-8")
        if push_transcripts and snapshot.transcript:
            files["transcript/raw_transcript.txt"] = (snapshot.transcript["original_text"].rstrip() + "\n").encode("utf-8")
            files["transcript/speaker_labelled_transcript.txt"] = _speaker_labelled_text(
                snapshot.transcript.get("speaker_labelled_path")
            ).encode("utf-8")
        return files

    @staticmethod
    def _docx(snapshot: ApprovedSessionSnapshot) -> bytes:
        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)
        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(11)
        for name in ("Title", "Heading 1", "Heading 2"):
            styles[name].font.name = "Arial"
            styles[name].font.color.rgb = RGBColor(0, 0, 0)
        title = document.add_paragraph(style="Title")
        title.add_run("Clinical Note")
        document.add_paragraph(
            f"Session {snapshot.session_id} | {snapshot.specialty} | {snapshot.session_date}"
        )
        for line in snapshot.note["content"].splitlines():
            stripped = line.strip()
            if not stripped:
                document.add_paragraph()
            elif stripped.startswith("# "):
                document.add_heading(stripped[2:], level=1)
            elif stripped.startswith("## "):
                document.add_heading(stripped[3:], level=2)
            elif stripped.startswith("- "):
                document.add_paragraph(stripped[2:], style="List Bullet")
            else:
                document.add_paragraph(stripped)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _pdf(snapshot: ApprovedSessionSnapshot) -> bytes:
        buffer = io.BytesIO()
        styles = getSampleStyleSheet()
        title = ParagraphStyle("ClinicalTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, alignment=TA_CENTER, textColor="#000000", spaceAfter=12)
        body = ParagraphStyle("ClinicalBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=6)
        h1 = ParagraphStyle("ClinicalH1", parent=styles["Heading1"], fontName="Helvetica-Bold", textColor="#000000", fontSize=14, leading=17, spaceBefore=10, spaceAfter=6)
        h2 = ParagraphStyle("ClinicalH2", parent=h1, fontSize=12, leading=15)
        story = [Paragraph("Clinical Note", title), Paragraph(html.escape(f"Session {snapshot.session_id} | {snapshot.specialty} | {snapshot.session_date}"), body), Spacer(1, 8)]
        for line in snapshot.note["content"].splitlines():
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 6))
            elif stripped.startswith("# "):
                story.append(Paragraph(html.escape(stripped[2:]), h1))
            elif stripped.startswith("## "):
                story.append(Paragraph(html.escape(stripped[3:]), h2))
            elif stripped.startswith("- "):
                story.append(Paragraph("&#8226; " + html.escape(stripped[2:]), body))
            else:
                story.append(Paragraph(html.escape(stripped), body))
        document = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=0.8 * inch, leftMargin=0.8 * inch, topMargin=0.75 * inch, bottomMargin=0.75 * inch, title="Clinical Note", author="Clinical Scribe")
        document.build(story)
        return buffer.getvalue()

    @staticmethod
    def _print_html(snapshot: ApprovedSessionSnapshot) -> str:
        note = html.escape(snapshot.note["plain_text"])
        metadata = html.escape(f"Session {snapshot.session_id} | {snapshot.specialty} | {snapshot.session_date}")
        return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Clinical Note</title>
<style>@page {{ size: Letter; margin: 0.75in; }} body {{ font: 11pt/1.45 Arial, sans-serif; color: #000; }} h1 {{ font-size: 20pt; }} .metadata {{ color: #333; }} pre {{ white-space: pre-wrap; font: inherit; }} @media screen {{ body {{ max-width: 7in; margin: 2rem auto; }} }}</style>
</head><body><h1>Clinical Note</h1><p class="metadata">{metadata}</p><pre>{note}</pre><script>window.addEventListener("load", () => window.print());</script></body></html>"""
